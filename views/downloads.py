import streamlit as st
import pandas as pd
import io
from docx import Document
from database import run_query

import os
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_docx(df, title, is_judge_sheet=False, prog_row=None):
    doc = Document()
    
    # Fetch event info
    try:
        df_info = run_query("SELECT * FROM event_info LIMIT 1")
    except Exception:
        df_info = pd.DataFrame()
        
    school_name = "Tournament"
    sports_name = "Annual Sports Day"
    year = ""
    logo_path = None
    
    if not df_info.empty:
        school_name = df_info['school_name'].iloc[0] if pd.notna(df_info['school_name'].iloc[0]) else school_name
        sports_name = df_info['sports_name'].iloc[0] if pd.notna(df_info['sports_name'].iloc[0]) else sports_name
        year = df_info['year'].iloc[0] if pd.notna(df_info['year'].iloc[0]) else year
        logo_path = df_info['logo_path'].iloc[0] if pd.notna(df_info['logo_path'].iloc[0]) else None

    # Header section
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    if logo_path and os.path.exists(logo_path):
        run = header_para.add_run()
        run.add_picture(logo_path, width=Inches(1.0))
        header_para.add_run("\n")
        
    run = header_para.add_run(f"{school_name}\n")
    run.font.size = Pt(16)
    run.bold = True
    
    if year:
        run = header_para.add_run(f"{sports_name} - {year}\n")
    else:
        run = header_para.add_run(f"{sports_name}\n")
    run.font.size = Pt(12)
    
    header_para.add_run("\n")
    run = header_para.add_run(title)
    run.font.size = Pt(14)
    run.bold = True
    
    if is_judge_sheet and prog_row is not None:
        pdate = prog_row.get('date', '__________')
        ptime = prog_row.get('time', '__________')
        pvenue = prog_row.get('venue', '__________')
        pdate = pdate if pd.notna(pdate) and pdate else '__________'
        ptime = ptime if pd.notna(ptime) and ptime else '__________'
        pvenue = pvenue if pd.notna(pvenue) and pvenue else '__________'
        
        details_para = doc.add_paragraph()
        details_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        d_run = details_para.add_run(f"Date: {pdate}    Time: {ptime}    Venue: {pvenue}")
        d_run.font.size = Pt(11)
        
        judge_para = doc.add_paragraph()
        judge_para.add_run("\nJudge's Name: _______________________          Signature: _______________________")
        judge_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add a table for data
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = 'Table Grid'
    
    # Add header row
    hdr_cells = table.rows[0].cells
    for i, col_name in enumerate(df.columns):
        hdr_cells[i].text = str(col_name)
    
    # Add data rows
    for index, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val) if pd.notna(val) else ""
            
    # Add Footer with signatures
    doc.add_paragraph("\n\n\n")
    
    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.autofit = True
    
    cell_left = sig_table.cell(0, 0)
    cell_left.text = "________________________\nSports Chairman"
    cell_left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    cell_center = sig_table.cell(0, 1)
    cell_center.text = "________________________\nSports Coordinator"
    cell_center.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    cell_right = sig_table.cell(0, 2)
    cell_right.text = "________________________\nPrincipal"
    cell_right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def get_winners_df(event_mode, prog_id=None):
    import re
    df_scores = run_query("SELECT * FROM scores WHERE position IN (1, 2, 3)")
    if df_scores.empty: return pd.DataFrame()
    
    if prog_id is not None:
        df_scores = df_scores[df_scores['program_id'] == prog_id]
        if df_scores.empty: return pd.DataFrame()
        
    def extract_chest(val):
        m = re.search(r'\(Chest:\s*(\d+)\)', str(val))
        return int(m.group(1)) if m else None
        
    df_scores['Extracted_Chest'] = df_scores['student_or_group_leader'].apply(extract_chest)
    
    # Query updated to fetch division natively
    df_students = run_query("SELECT chest_number, name, class_name, division, category, group_name FROM students")
    df_merged = pd.merge(df_scores, df_students, left_on='Extracted_Chest', right_on='chest_number', how='left', suffixes=('_score', '_student'))
    
    df_programs = run_query("SELECT id, name, category as prog_category, gender, type, stage_type FROM programs WHERE event_mode = ?", (event_mode,))
    df_merged = pd.merge(df_merged, df_programs, left_on='program_id', right_on='id', how='left', suffixes=('', '_prog'))
    
    # Drop rows where there's no program match (because of event_mode exclusion)
    df_merged = df_merged.dropna(subset=['id'])
    
    def get_clean_name(row):
        if pd.notna(row.get('name_student')): return row['name_student']
        s = str(row['student_or_group_leader'])
        return re.sub(r'\s*\(Chest:.*\)', '', s)
        
    df_merged['Clean_Name'] = df_merged.apply(get_clean_name, axis=1)
    
    df_out = pd.DataFrame()
    df_out['Program Name'] = df_merged['name_prog'].astype(str)
    df_out['Gender'] = df_merged['gender'].fillna('Unknown')
    df_out['Category'] = df_merged['prog_category'].fillna('Unknown')
    df_out['Type'] = df_merged['type'].fillna('Unknown')
    df_out['Stage Type'] = df_merged['stage_type'].fillna('')
    df_out['Position'] = df_merged['position'].apply(lambda x: f"{int(x)}st" if x==1 else (f"{int(x)}nd" if x==2 else (f"{int(x)}rd" if x==3 else str(x))))
    df_out['Chest Number'] = df_merged['Extracted_Chest'].fillna('')
    df_out['Student Name'] = df_merged['Clean_Name']
    df_out['Group Name'] = df_merged['group_name_score']
    df_out['Class'] = df_merged['class_name'].fillna('')
    try:
        df_out['Division'] = df_merged['division'].fillna('')
    except KeyError:
        df_out['Division'] = ''
    
    df_out['Chest Number'] = df_out['Chest Number'].apply(lambda x: str(int(x)) if x != '' else '')
    
    return df_out.sort_values(by=['Category', 'Program Name', 'Position'])

def get_checklist_df(event_mode, sort_mode='Category'):
    df_prog = run_query("SELECT date, time, category, name, type, gender, venue FROM programs WHERE event_mode = ?", (event_mode,))
    if df_prog.empty: return pd.DataFrame()
    
    df_out = pd.DataFrame()
    df_out['Date'] = df_prog['date'].fillna('TBD')
    df_out['Time'] = df_prog['time'].fillna('TBD')
    df_out['Category'] = df_prog['category'].fillna('Unknown')
    df_out['Program Name'] = df_prog['name'].astype(str)
    df_out['Format'] = df_prog['gender'].astype(str) + " - " + df_prog['type'].astype(str)
    df_out['Venue'] = df_prog['venue'].fillna('')
    df_out['[  ] Event Over'] = ''
    df_out['[  ] Marks Uploaded'] = ''
    
    if sort_mode == 'Chronological':
        return df_out.sort_values(by=['Date', 'Time', 'Category', 'Program Name'])
    else:
        return df_out.sort_values(by=['Category', 'Date', 'Time', 'Program Name'])

def generate_winners_docx(df_filtered):
    doc = Document()
    
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_para.add_run("Tournament\n")
    run.font.size = Pt(16)
    run.bold = True
    run = header_para.add_run("Annual Sports Day\n\n")
    run.font.size = Pt(12)
    run = header_para.add_run("Winners List\n")
    run.font.size = Pt(14)
    run.bold = True
    
    doc.add_paragraph()

    req_cols = ["Chest Number", "Student Name", "Group Name", "Class", "Division", "Position"]
    
    if not df_filtered.empty:
        for cat_name, cat_df in df_filtered.groupby('Category'):
            doc.add_heading(str(cat_name), level=1)
            
            # Sort prog_df groupings safely
            for (prog_name, prog_type, prog_gender), prog_df in cat_df.groupby(['Program Name', 'Type', 'Gender']):
                doc.add_heading(f"{prog_name} ({prog_gender}) - {prog_type}", level=2)
                
                table_df = prog_df[req_cols]
                
                table = doc.add_table(rows=1, cols=len(table_df.columns))
                table.style = 'Table Grid'
                
                hdr_cells = table.rows[0].cells
                for i, col_name in enumerate(table_df.columns):
                    hdr_cells[i].text = str(col_name)
                    hdr_cells[i].paragraphs[0].runs[0].bold = True
                
                for _, row in table_df.iterrows():
                    row_cells = table.add_row().cells
                    for i, val in enumerate(row):
                        row_cells[i].text = str(val) if pd.notna(val) else ""
                
                doc.add_paragraph()
                
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def render():
    current_mode = st.session_state.get('event_mode', 'Sports')
    st.title(f"🖨️ Downloads & Prints ({current_mode})")
    st.markdown("Generate and download specific sheets for the sports event, such as Judges' Scoring Sheets and Groupwise Participant Lists.")

    tab1, tab2, tab3, tab4 = st.tabs(["📝 Judges' Scoring Sheets", "🏆 Groupwise Final Lists", "🏅 Winners List", "📋 Coordinator Checklist"])

    with tab1:
        st.subheader("Judges' Scoring Sheet")
        
        # 1. Provide a way to filter by Category & Gender
        try:
            df_categories = run_query("SELECT name FROM categories")
            categories = df_categories['name'].tolist() if not df_categories.empty else []
            df_genders = run_query("SELECT name FROM genders")
            genders = df_genders['name'].tolist() if not df_genders.empty else []
            genders_opts = ["All"] + genders + ["Mixed"] if genders else ["All", "Mixed"]
        except Exception:
            categories = []
            genders_opts = ["All", "Mixed"]

        if not categories:
            st.warning("No categories found. Please add them in Settings or Upload Data.")
        else:
            col_f1, col_f2, col_f3 = st.columns(3)
            with col_f1:
                selected_gender = st.selectbox("Select Gender", genders_opts, key="judge_gen")
            with col_f2:
                selected_category = st.selectbox("Select Category", categories, key="judge_cat")
            with col_f3:
                selected_type = st.selectbox("Select Event Type", ["All", "Individual", "Group"], key="judge_type")

            # 2. Filter programs by the selected category and gender
            try:
                query_prog = "SELECT id, name, type, category, gender, stage_type, date, time, venue FROM programs WHERE category = ? AND event_mode = ?"
                params_prog = [selected_category, current_mode]
                
                if selected_gender != "All":
                    query_prog += " AND gender = ?"
                    params_prog.append(selected_gender)
                    
                if selected_type != "All":
                    query_prog += " AND type = ?"
                    params_prog.append(selected_type)
                    
                df_programs = run_query(query_prog, tuple(params_prog))
            except Exception:
                df_programs = pd.DataFrame()

            if df_programs.empty:
                st.info(f"No {current_mode} programs found for category: {selected_category}")
            else:
                program_display = [f"{row['name']} ({row['gender']} - {row['type']}{' - ' + str(row['stage_type']) if pd.notna(row['stage_type']) else ''})" for _, row in df_programs.iterrows()]
                selected_prog_str = st.selectbox("Select Program", program_display, key="judge_prog")
                
                # Get the selected program details
                prog_idx = program_display.index(selected_prog_str)
                prog_row = df_programs.iloc[prog_idx]

                time_option = st.radio("Schedule Information", ["🗄️ Use Database Schedule", "✍️ Enter Manually"])
                
                manual_date = ""
                manual_time = ""
                manual_venue = ""
                if time_option == "✍️ Enter Manually":
                    col_m1, col_m2, col_m3 = st.columns(3)
                    with col_m1: manual_date = st.text_input("Date")
                    with col_m2: manual_time = st.text_input("Time")
                    with col_m3: manual_venue = st.text_input("Venue")
                
                if st.button("Generate Scoring Sheet", key="btn_gen_judge"):
                    prog_dict = prog_row.to_dict()
                    if time_option == "✍️ Enter Manually":
                        prog_dict['date'] = manual_date
                        prog_dict['time'] = manual_time
                        prog_dict['venue'] = manual_venue
                        
                    if prog_dict['type'] == 'Individual':
                        query = """
                        SELECT 
                            s.chest_number as 'Chest Number', 
                            s.name as 'Student Name', 
                            s.group_name as 'Group Name' 
                        FROM individual_participants ip
                        JOIN students s ON ip.student_id = s.id
                        WHERE ip.program_id = ?
                        ORDER BY s.chest_number ASC
                        """
                        df_participants = run_query(query, (int(prog_dict['id']),))
                        
                        if df_participants.empty:
                            st.warning("No students have been assigned to this Individual program. Please assign them in 'Upload & Manage Data -> Individual Assignation'.")
                        else:
                            df_participants.insert(0, 'Program Name', prog_dict['name'])
                            df_participants['Position/Score'] = ''
                            df_participants['Remarks'] = ''
                            
                            st.success("Scoring sheet generated!")
                            st.dataframe(df_participants, use_container_width=True, hide_index=True)
                            
                            csv = df_participants.to_csv(index=False).encode('utf-8')
                            docx_bytes = generate_docx(df_participants, f"Judges Scoring Sheet: {prog_dict['name']} ({selected_category})", is_judge_sheet=True, prog_row=prog_dict)
                            
                            col1, col2 = st.columns(2)
                            with col1:
                                st.download_button(
                                    label="📥 Download CSV",
                                    data=csv,
                                    file_name=f"Judges_Sheet_{prog_dict['name']}_{selected_category}.csv",
                                    mime='text/csv',
                                    key="dl_csv_ind"
                                )
                            with col2:
                                st.download_button(
                                    label="📝 Download Word File (.docx)",
                                    data=docx_bytes,
                                    file_name=f"Judges_Sheet_{prog_dict['name']}_{selected_category}.docx",
                                    mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                                    key="dl_docx_ind"
                                )
                        
                    else:
                        query_grp = """
                        SELECT 
                            s.chest_number as 'Chest Number', 
                            s.name as 'Student Name (Leader)', 
                            s.group_name as 'Group Name',
                            GROUP_CONCAT(gp.participant_name, ', ') as 'Participants'
                        FROM group_participants gp
                        JOIN students s ON gp.leader_chest_number = s.chest_number
                        WHERE gp.program_id = ?
                        GROUP BY s.chest_number, s.name, s.group_name
                        ORDER BY s.chest_number ASC
                        """
                        df_groups = run_query(query_grp, (int(prog_dict['id']),))
                        if df_groups.empty:
                            st.warning("No registered teams found for this group program. Please construct teams in 'Group Team Setup'.")
                        else:
                            df_groups.insert(0, 'Program Name', prog_dict['name'])
                            df_groups['Position/Score'] = ''
                            df_groups['Remarks'] = ''
                            
                            st.success("Scoring sheet generated!")
                            st.dataframe(df_groups, use_container_width=True, hide_index=True)
                            
                            csv = df_groups.to_csv(index=False).encode('utf-8')
                            docx_bytes = generate_docx(df_groups, f"Judges Scoring Sheet: {prog_dict['name']} ({selected_category})", is_judge_sheet=True, prog_row=prog_dict)

                            
                            col1, col2 = st.columns(2)
                            with col1:
                                st.download_button(
                                    label="📥 Download CSV",
                                    data=csv,
                                    file_name=f"Judges_Sheet_Group_{prog_dict['name']}_{selected_category}.csv",
                                    mime='text/csv',
                                    key="dl_csv_grp"
                                )
                            with col2:
                                st.download_button(
                                    label="📝 Download Word File (.docx)",
                                    data=docx_bytes,
                                    file_name=f"Judges_Sheet_Group_{prog_dict['name']}_{selected_category}.docx",
                                    mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                                    key="dl_docx_grp"
                                )

    with tab2:
        st.subheader("Final lists of Groupwise Participants")
        st.markdown("Generate a list of all students and their participating programs for a specific group, perfect for distributing to group leaders.")
        
        try:
            df_grp_list = run_query("SELECT name FROM groups")
            groups_list = df_grp_list['name'].tolist() if not df_grp_list.empty else []
        except Exception:
            groups_list = []

        if not groups_list:
            st.warning("No groups found. Please add them in Settings or Upload Data.")
        else:
            selected_group = st.selectbox("Select Group", groups_list, key="final_list_grp")
            selected_student_gender = st.selectbox("Select Gender Filter", ["All", "Boys", "Girls"], key="final_list_gender")
            
            if st.button("Generate Groupwise List", key="btn_gen_grouplist"):
                
                params = [selected_group]
                filter_cond = ""
                if selected_student_gender != "All":
                    filter_cond = " AND s.gender = ?"
                    params.append(selected_student_gender)

                union_params = params + params

                query_standard = f"""
                SELECT * FROM (
                    SELECT 
                        s.category as 'Category', 
                        s.chest_number as 'Chest Number', 
                        s.name as 'Student Name', 
                        s.gender as 'Gender',
                        p.name as 'Program Name'
                    FROM individual_participants ip
                    JOIN students s ON ip.student_id = s.id
                    JOIN programs p ON ip.program_id = p.id
                    WHERE s.group_name = ? AND p.type = 'Individual' AND p.event_mode = ? {filter_cond}
                    UNION
                    SELECT 
                        s.category as 'Category', 
                        s.chest_number as 'Chest Number', 
                        s.name as 'Student Name', 
                        s.gender as 'Gender',
                        p.name || ' (Group)' as 'Program Name'
                    FROM students s
                    JOIN group_participants gp ON s.chest_number = gp.leader_chest_number
                    JOIN programs p ON gp.program_id = p.id
                    WHERE s.group_name = ? AND p.event_mode = ? {filter_cond}
                ) ORDER BY Category ASC, "Chest Number" ASC, "Program Name" ASC
                """
                
                query_scheduled = f"""
                SELECT * FROM (
                    SELECT 
                        s.category as 'Category', 
                        s.chest_number as 'Chest Number', 
                        s.name as 'Student Name', 
                        s.gender as 'Gender',
                        p.name as 'Program Name',
                        p.date as 'Date',
                        p.time as 'Time',
                        p.venue as 'Venue'
                    FROM individual_participants ip
                    JOIN students s ON ip.student_id = s.id
                    JOIN programs p ON ip.program_id = p.id
                    WHERE s.group_name = ? AND p.type = 'Individual' AND p.event_mode = ? {filter_cond}
                    UNION
                    SELECT 
                        s.category as 'Category', 
                        s.chest_number as 'Chest Number', 
                        s.name as 'Student Name', 
                        s.gender as 'Gender',
                        p.name || ' (Group)' as 'Program Name',
                        p.date as 'Date',
                        p.time as 'Time',
                        p.venue as 'Venue'
                    FROM students s
                    JOIN group_participants gp ON s.chest_number = gp.leader_chest_number
                    JOIN programs p ON gp.program_id = p.id
                    WHERE s.group_name = ? AND p.event_mode = ? {filter_cond}
                ) ORDER BY Date ASC, Time ASC, Category ASC, "Chest Number" ASC, "Program Name" ASC
                """
                
                import collections
                if isinstance(union_params, tuple): union_params = list(union_params)
                
                # We need to insert `current_mode` into the params at the correct positions!
                # params was: [selected_group] or [selected_group, selected_student_gender]
                
                bind_params1 = [selected_group, current_mode]
                if selected_student_gender != "All": bind_params1.append(selected_student_gender)
                
                bind_params_total = bind_params1 + bind_params1
                
                df_standard = run_query(query_standard, tuple(bind_params_total))
                df_scheduled = run_query(query_scheduled, tuple(bind_params_total))
                
                if df_standard.empty:
                    st.info(f"No participating students or programs found for {selected_group} ({selected_student_gender}).")
                else:
                    st.success(f"Groupwise lists generated for {selected_group} ({selected_student_gender})!")
                    
                    st.markdown("### Preview (With Schedule)")
                    st.dataframe(df_scheduled, use_container_width=True, hide_index=True)
                    
                    st.markdown("---")
                    st.markdown("#### Option 1: Standard Groupwise List (No Schedule)")
                    csv_1 = df_standard.to_csv(index=False).encode('utf-8')
                    docx_1 = generate_docx(df_standard, f"Final Participation List: {selected_group} ({selected_student_gender})")
                    c1, c2 = st.columns(2)
                    c1.download_button("📥 Download CSV (Standard)", csv_1, f"Group_List_Std_{selected_group}_{selected_student_gender}.csv", "text/csv", key="btn_csv_std")
                    c2.download_button("📝 Download Word (Standard)", docx_1, f"Group_List_Std_{selected_group}_{selected_student_gender}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="btn_docx_std")
                    
                    st.markdown("#### Option 2: Groupwise List WITH Time Schedule")
                    csv_2 = df_scheduled.to_csv(index=False).encode('utf-8')
                    docx_2 = generate_docx(df_scheduled, f"Scheduled Participation List: {selected_group} ({selected_student_gender})")
                    c3, c4 = st.columns(2)
                    c3.download_button("📥 Download CSV (Scheduled)", csv_2, f"Group_List_Sched_{selected_group}_{selected_student_gender}.csv", "text/csv", key="btn_csv_sch")
                    c4.download_button("📝 Download Word (Scheduled)", docx_2, f"Group_List_Sched_{selected_group}_{selected_student_gender}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="btn_docx_sch")

    with tab3:
        st.subheader("🏅 Final Winners Documentation")
        st.markdown("Use the filters below to dynamically fetch exact lists of 1st, 2nd, and 3rd place winners to print in an elegant, structured Word package.")
        
        # UI Filters
        try:
            db_cats = run_query("SELECT name FROM categories")['name'].tolist()
            db_gens = run_query("SELECT name FROM genders")['name'].tolist()
            db_gens.append("Mixed")
        except Exception:
            db_cats = []
            db_gens = ["Boys", "Girls", "Mixed"]
            
        c1, c2, c3 = st.columns(3)
        with c1:
            w_cat = st.selectbox("Category", ["ALL"] + db_cats, key="win_flt_cat")
        with c2:
            w_type = st.selectbox("Event Type", ["ALL", "Individual", "Group"], key="win_flt_type")
        with c3:
            w_gen = st.selectbox("Gender", ["ALL"] + db_gens, key="win_flt_gen")
            
        w_stage = "ALL"
        if current_mode == 'Arts':
            c4, c5 = st.columns(2)
            with c4:
                w_stage = st.selectbox("Stage Type (Arts Only)", ["ALL", "On-stage", "Off-stage"], key="win_flt_stage")
                
        st.markdown("---")
        
        df_all_winners = get_winners_df(event_mode=current_mode, prog_id=None)
        
        if df_all_winners.empty:
            st.info("No winners have been logged in the entire system yet.")
        else:
            # Apply Filters
            if w_cat != "ALL":
                df_all_winners = df_all_winners[df_all_winners["Category"] == w_cat]
            if w_type != "ALL":
                df_all_winners = df_all_winners[df_all_winners["Type"] == w_type]
            if w_gen != "ALL":
                df_all_winners = df_all_winners[df_all_winners["Gender"] == w_gen]
            if w_stage != "ALL":
                df_all_winners = df_all_winners[df_all_winners["Stage Type"] == w_stage]
                
            if df_all_winners.empty:
                st.warning("No winners logged corresponding to these specific filters.")
            else:
                st.success(f"Found {len(df_all_winners)} matching winners across events!")
                
                # Show Preview
                display_df = df_all_winners[["Program Name", "Category", "Gender", "Type", "Position", "Chest Number", "Student Name", "Group Name", "Class", "Division"]]
                if current_mode == 'Arts':
                    display_df.insert(4, 'Stage Type', df_all_winners['Stage Type'])
                    
                st.dataframe(display_df, use_container_width=True, hide_index=True)
                
                csv_m = df_all_winners.to_csv(index=False).encode('utf-8')
                docx_m = generate_winners_docx(df_all_winners)
                
                co1, co2 = st.columns(2)
                co1.download_button("📥 Download Filtered Winners (.csv)", csv_m, "Winners_List_Export.csv", "text/csv", key="btn_w_csv_m", type="primary")
                co2.download_button("📝 Download Filtered Master Book (.docx)", docx_m, "Winners_List_Export.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="btn_w_docx_m", type="primary")

    with tab4:
        st.subheader("📋 Coordinator Operations Checklist")
        st.markdown("Generate a master paper checklist meant directly for the central Coordinator. Contains intentionally blank columns designed specifically to be pen-ticked live as events finalize and scores get uploaded.")
        
        st.markdown("---")
        mode = st.radio("Select Print Sorting Mode:", ["📂 Option 1: Category Wise (Grouped by Event Category)", "⏱️ Option 2: Time and Date Wise (Strict Chronological Format)"])
        
        if st.button("Generate Operator Checklist 📄", type="primary"):
            sort_val = 'Chronological' if 'Time' in mode else 'Category'
            df_check = get_checklist_df(event_mode=current_mode, sort_mode=sort_val)
            
            if df_check.empty:
                st.warning("No programs have been registered in the system.")
            else:
                st.success("Checklist generated successfully!")
                st.dataframe(df_check, use_container_width=True, hide_index=True)
                
                csv_c = df_check.to_csv(index=False).encode('utf-8')
                docx_c = generate_docx(df_check, f"Coordinator Event Checklist ({sort_val} Sort)")
                
                cx1, cx2 = st.columns(2)
                cx1.download_button("📥 Download Checklist (.csv)", csv_c, f"Checklist_{sort_val}.csv", "text/csv", key="btn_c_csv")
                cx2.download_button("📝 Download Checklist Word (.docx)", docx_c, f"Checklist_{sort_val}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key="btn_c_docx")
