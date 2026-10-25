import streamlit as st
import pandas as pd
from database import run_query, transaction
import io
from docx import Document

from docx.enum.text import WD_ALIGN_PARAGRAPH
from views.downloads import add_doc_header
from views.ui import hero, require_admin


def generate_schedule_docx(df):
    doc = Document()
    add_doc_header(doc, "Event Time Schedule")

    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    for i, col_name in enumerate(df.columns):
        hdr_cells[i].text = str(col_name)

    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val) if pd.notna(val) else ""

    # Add Footer with signatures
    doc.add_paragraph("\n\n\n")

    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.autofit = True
    for i, (label, align) in enumerate([("Sports Chairman", WD_ALIGN_PARAGRAPH.LEFT),
                                         ("Sports Coordinator", WD_ALIGN_PARAGRAPH.CENTER),
                                         ("Principal", WD_ALIGN_PARAGRAPH.RIGHT)]):
        cell = sig_table.cell(0, i)
        cell.text = f"________________________\n{label}"
        cell.paragraphs[0].alignment = align

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _parse(series, kind):
    """Turn stored text into date/time objects. Returns (parsed, unparseable_mask)."""
    text = series.fillna("").astype(str).str.strip()
    text = text.where(text != "")
    # Values saved by the app are ISO (YYYY-MM-DD / HH:MM); anything older is typed text, read day-first (Indian format)
    iso_fmt = "%Y-%m-%d" if kind == "date" else "%H:%M"
    parsed = pd.to_datetime(text, errors="coerce", format=iso_fmt)
    fallback = pd.to_datetime(text, errors="coerce", format="mixed", dayfirst=True)
    parsed = parsed.fillna(fallback)
    text = text.fillna("")
    bad = (text != "") & parsed.isna()
    values = parsed.dt.date if kind == "date" else parsed.dt.time
    return values.where(parsed.notna(), None), bad


def render():
    current_mode = st.session_state.get('event_mode', 'Sports')
    hero(f"{current_mode} Event Schedule Planner", "Set the date, time and venue of every program, then download the final schedule.", "🗓️")

    df_prog = run_query(
        "SELECT id, name, category, gender, type, date, time, venue FROM programs WHERE event_mode = ? ORDER BY date ASC, time ASC, name ASC",
        (current_mode,),
    )

    if df_prog.empty:
        st.info("No programs found. Please add programs in 'Upload & Manage Data' first.")
        return

    if require_admin():
        st.subheader("Manage Program Schedule")
        st.caption("Pick the Date and Time with the calendar/clock pickers and type the Venue, then click Save.")

        df_edit = df_prog.copy()
        df_edit['date'], bad_dates = _parse(df_prog['date'], "date")
        df_edit['time'], bad_times = _parse(df_prog['time'], "time")
        bad = bad_dates | bad_times
        if bad.any():
            st.warning("These programs have a date/time the app couldn't read. Please re-enter them with the pickers "
                       "(the old text is shown here and will be replaced when you save):")
            st.dataframe(df_prog.loc[bad, ['name', 'category', 'date', 'time']], hide_index=True, width="stretch")

        edited_prog = st.data_editor(
            df_edit,
            hide_index=True,
            width="stretch",
            disabled=["id", "name", "category", "gender", "type"],
            column_config={
                "id": None,
                "name": "Program",
                "category": "Category",
                "gender": "Gender",
                "type": "Type",
                "date": st.column_config.DateColumn("Date", format="DD/MM/YYYY"),
                "time": st.column_config.TimeColumn("Time", format="hh:mm a", step=300),
                "venue": st.column_config.TextColumn("Venue"),
            },
            key="schedule_editor",
        )

        # Clash detection: same venue, same date and time
        clash = edited_prog.dropna(subset=['date', 'time'])
        clash = clash[clash['venue'].fillna("").str.strip() != ""]
        clash = clash[clash.duplicated(subset=['date', 'time', 'venue'], keep=False)]
        if not clash.empty:
            st.error("⚠️ Venue clash — these programs are at the same venue at the same time:")
            st.dataframe(clash[['name', 'category', 'date', 'time', 'venue']], hide_index=True, width="stretch")

        if st.button("💾 Save Schedule Settings", type="primary"):
            with transaction() as conn:
                for _, row in edited_prog.iterrows():
                    d = row['date'].isoformat() if pd.notna(row['date']) and row['date'] is not None else None
                    t = row['time'].strftime("%H:%M") if pd.notna(row['time']) and row['time'] is not None else None
                    v = row['venue'] if pd.notna(row['venue']) and str(row['venue']).strip() else None
                    conn.execute("UPDATE programs SET date = ?, time = ?, venue = ? WHERE id = ?", (d, t, v, int(row['id'])))
            st.toast("Schedule saved!", icon="✅")
            st.rerun()

    st.markdown("---")
    st.subheader("Final Schedule")

    view_df = run_query(
        "SELECT date as Date, time as Time, venue as Venue, category as Category, name as Program, gender as Gender, type as Type "
        "FROM programs WHERE event_mode = ? ORDER BY date IS NULL, date ASC, time ASC, venue ASC",
        (current_mode,),
    )
    dates = pd.to_datetime(view_df['Date'], errors="coerce", format="mixed")
    times = pd.to_datetime(view_df['Time'], errors="coerce", format="mixed")
    view_df['Date'] = dates.dt.strftime("%d %b %Y").where(dates.notna(), view_df['Date'])
    view_df['Time'] = times.dt.strftime("%I:%M %p").where(times.notna(), view_df['Time'])

    st.dataframe(view_df, width="stretch", hide_index=True)

    col1, col2 = st.columns(2)
    with col1:
        st.download_button("📥 Download Schedule as CSV", view_df.to_csv(index=False).encode('utf-8'),
                           "Event_Schedule.csv", 'text/csv', key="dl_sched_csv")
    with col2:
        st.download_button("📝 Download Schedule as Word File", generate_schedule_docx(view_df), "Event_Schedule.docx",
                           'application/vnd.openxmlformats-officedocument.wordprocessingml.document', key="dl_sched_docx")
